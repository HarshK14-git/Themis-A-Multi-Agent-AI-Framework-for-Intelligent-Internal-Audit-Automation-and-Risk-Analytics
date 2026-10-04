"""
Evidence Extraction (Phase 12 extension)
Reads text out of any evidence file type. Images go through OCR (pixels ->
text). Spreadsheets and slide decks already have machine-readable text, so
they're read directly - running OCR on them would be strictly worse.
"""
import os
from pypdf import PdfReader
from openpyxl import load_workbook
from pptx import Presentation

_ocr_reader = None

def get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        import easyocr
        _ocr_reader = easyocr.Reader(["en"], gpu=False)
    return _ocr_reader

def extract_from_image(path):
    reader = get_ocr_reader()
    results = reader.readtext(path, detail=0)
    return "\n".join(results)

def extract_from_xlsx(path):
    wb = load_workbook(path, read_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join(str(c) for c in row if c is not None))
    return "\n".join(lines)

def extract_from_pptx(path):
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = "\t".join(cell.text for cell in row.cells)
                    lines.append(row_text)
    return "\n".join(lines)

def extract_from_pdf(path):
    reader = PdfReader(path)
    lines = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text()
        if text and text.strip():
            lines.append(f"--- Page {i} ---")
            lines.append(text)
        else:
            lines.append(f"--- Page {i}: no extractable text (likely a scanned image - would need OCR) ---")
    return "\n".join(lines)

def extract_evidence_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".png", ".jpg", ".jpeg"):
        return extract_from_image(path)
    elif ext == ".xlsx":
        return extract_from_xlsx(path)
    elif ext == ".pptx":
        return extract_from_pptx(path)
    elif ext == ".pdf":
        return extract_from_pdf(path)
    else:
        return f"[Unsupported file type: {ext}]"


if __name__ == "__main__":
    folder = "data/raw"
    for filename in os.listdir(folder):
        path = os.path.join(folder, filename)
        print(f"\n{'='*60}\n{filename}\n{'='*60}")
        text = extract_evidence_text(path)
        print(text[:800])